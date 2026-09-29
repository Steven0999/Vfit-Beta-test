#!/usr/bin/env python3
"""Rebuild VFIT's owned UK catalogue from the official CoFID 2021 workbook.

Usage: python3 tools/import-food-catalog.py /path/to/cofid-2021.xlsx
Requires openpyxl. No API key or live nutrition service is used by the app.
"""
import hashlib
import json
import math
import pathlib
import sys
from collections import Counter

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
SOURCE_URL = 'https://assets.publishing.service.gov.uk/media/60538b91e90e07527df82ae4/McCance_Widdowsons_Composition_of_Foods_Integrated_Dataset_2021..xlsx'
COLUMNS = ['key', 'code', 'name', 'category', 'basisUnit', 'calories', 'protein', 'carbs', 'fat', 'fiber', 'sugar', 'satFat', 'sodiumMg', 'cholesterol', 'foodGroup', 'description', 'reference', 'footnote', 'traceNutrients']


def value(raw):
    if str(raw).strip().lower() == 'tr':
        return 0
    try:
        number = float(raw)
        return number if math.isfinite(number) and number >= 0 else None
    except (TypeError, ValueError):
        return None


def main():
    path = pathlib.Path(sys.argv[1])
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # The published workbook uses 13-669 for two different foods. Join by name
    # AND code and give duplicate-code records stable separate storage keys.
    minerals = {(str(row[0]), str(row[1])): row for row in workbook['1.4 Inorganics'].iter_rows(min_row=4, values_only=True) if row[0]}
    rows = list(workbook['1.3 Proximates'].iter_rows(min_row=4, values_only=True))
    codes = Counter(str(row[0]) for row in rows if row[0] and row[1])
    categories = {'A': 'carbohydrate', 'B': 'dairy', 'C': 'protein', 'D': 'fruit-veg', 'F': 'fruit-veg', 'J': 'protein', 'M': 'protein', 'P': 'drink', 'Q': 'drink', 'S': 'meal', 'W': 'drink'}
    foods, excluded = [], []
    count = 0
    for row in rows:
        if not row[0] or not row[1]:
            continue
        count += 1
        raw = dict(zip(['calories', 'protein', 'carbs', 'fat', 'fiber', 'sugar', 'satFat', 'cholesterol'], [row[i] for i in [12, 9, 11, 10, 25, 16, 27, 46]]))
        mineral = minerals.get((str(row[0]), str(row[1])))
        raw['sodiumMg'] = mineral[7] if mineral else None
        nutrients = {key: value(item) for key, item in raw.items()}
        missing = [key for key in ['calories', 'protein', 'carbs', 'fat'] if nutrients[key] is None]
        if missing:
            excluded.append({'code': row[0], 'name': row[1], 'missing': missing})
            continue
        group = str(row[3] or '')
        code = str(row[0])
        key = code if codes[code] == 1 else code + '-' + hashlib.sha256(str(row[1]).encode()).hexdigest()[:8]
        record = {'key': key, 'code': code, 'name': str(row[1]), 'category': categories.get(group[:1], 'general'), 'basisUnit': 'ml' if group.startswith('Q') else 'g', 'foodGroup': group, 'description': str(row[2] or ''), 'reference': str(row[5] or ''), 'footnote': str(row[6] or ''), 'traceNutrients': [key for key, item in raw.items() if str(item).strip().lower() == 'tr'], **nutrients}
        foods.append([record[column] for column in COLUMNS])
    assert len({row[0] for row in foods}) == len(foods)
    result = {'schemaVersion': 1, 'version': 'cofid-2021-v1', 'source': {'name': 'Public Health England, UK CoFID 2021', 'url': 'https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid', 'downloadUrl': SOURCE_URL, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'license': 'Open Government Licence v3.0', 'licenseUrl': 'https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/', 'attribution': 'Contains public sector information licensed under the Open Government Licence v3.0. Source: Public Health England, McCance and Widdowson’s Composition of Foods Integrated Dataset (2021).'}, 'sourceFoodCount': count, 'foodCount': len(foods), 'excludedIncompleteFoods': excluded, 'columns': COLUMNS, 'foods': foods}
    output = ROOT / 'nutrition/data/uk-foods-2021.json'
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    print(f'{len(foods):,} complete UK foods; {len(excluded)} incomplete records excluded; {output.stat().st_size:,} bytes')


if __name__ == '__main__':
    main()
